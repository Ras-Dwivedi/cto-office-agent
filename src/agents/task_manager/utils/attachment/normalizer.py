import fitz           # PyMuPDF
import docx
import openpyxl


def normalize_text(path: str) -> list[str]:
    if path.endswith(".pdf"):
        return _from_pdf(path)
    if path.endswith(".docx"):
        return _from_docx(path)
    if path.endswith(".xlsx"):
        return _from_xlsx(path)
    return []


def _from_pdf(path):
    try:
        doc = fitz.open(path)
        lines = []
        for page in doc:
            lines.extend(page.get_text().split("\n"))
        return [l.strip() for l in lines if l.strip()]
    except Exception as e:
        return []

def _from_docx(path):
    try:
        doc = docx.Document(path)
        return [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    except Exception as e:
        return []

def _from_xlsx(path):
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(sheet.title)
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value:
                        lines.append(str(cell.value).strip())
        return lines
    except Exception as e:
        return []
